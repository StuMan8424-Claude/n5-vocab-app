#!/usr/bin/env python3
"""
Build script for the N5 Vocabulary Practice app.

Regenerates index.html (a single, self-contained, offline-capable HTML file)
by embedding the current vocabulary list into template.html.

By default it PULLS THE LIVE SHEET from the private tutoring repo first, so a
build is never made from a stale snapshot. The token is read from the
GITHUB_TOKEN environment variable (never stored in this public repo):

    GITHUB_TOKEN=<token> python3 build.py          # pull latest sheet, rebuild
    python3 build.py --no-pull                      # use the local .xlsx as-is
    python3 build.py --xlsx some_sheet.xlsx         # build from a specific file
    python3 build.py --json vocab_data.json         # build from an extracted export

Without a token and without --no-pull / --xlsx it falls back to the local
snapshot and says so loudly.

After building it prints what changed versus the previous vocab_data.json:
new / removed words, words newly marked learned, edited meanings.

Word ids are STABLE KEYS (kanji|reading, with the English meaning appended
only where two entries collide), not row numbers, so progress saved on a
device survives the sheet being re-sorted or having words inserted.
legacy_ids.json (the old row-number order) is embedded so devices that saved
progress under the old numeric ids are converted once on first load.

Requires: openpyxl (only when reading an .xlsx)
    pip install openpyxl --break-system-packages
"""
import argparse
import json
import os
import re
import sys
import urllib.request
from collections import Counter
from datetime import datetime, timezone

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(HERE, "template.html")
OUTPUT_PATH = os.path.join(HERE, "index.html")
DEFAULT_XLSX = os.path.join(HERE, "japanese_n5_vocabulary_updated.xlsx")
DEFAULT_JSON = os.path.join(HERE, "vocab_data.json")
LEGACY_IDS = os.path.join(HERE, "legacy_ids.json")

DEFAULT_REPO = "StuMan8424-Claude/japanese-n5-tutor"
DEFAULT_SHEET_PATH = "japanese_n5_vocabulary_updated.xlsx"

EXPECTED_HEADER = ["Kanji", "Hiragana", "Romaji", "English", "Category", "Word Type", "Learned", "Notes", "Last Used"]


# ---------- English sanitising ----------
# Some English meanings carry Japanese: grammar hints like
# "It would be better to ~ (た/ない form + ほうがいい)" or a list of example
# readings. As a flash-card prompt or a multiple-choice option that hands over
# the very answer being asked for. The clean English goes in `e` (prompts,
# options, answer matching); what was stripped is kept in `n` (a note) so it
# can still be shown where it's a help rather than a spoiler.
JP_RE = re.compile(r"[\u3040-\u30ff\u3400-\u9fff\uff66-\uff9f]")


def _strip_jp_parentheticals(text):
    notes, out, i = [], [], 0
    while i < len(text):
        if text[i] == "(":
            depth, j = 0, i
            while j < len(text):
                if text[j] == "(":
                    depth += 1
                elif text[j] == ")":
                    depth -= 1
                    if depth == 0:
                        break
                j += 1
            # balanced match, so nested (...) like "は(particle)" is handled
            if j < len(text) and JP_RE.search(text[i + 1:j]):
                notes.append(text[i + 1:j].strip())
                while out and out[-1] == " ":
                    out.pop()
                i = j + 1
                continue
        out.append(text[i])
        i += 1
    return "".join(out), notes


def sanitize_english(english):
    # Return (clean_english, note). Idempotent; clean English never contains Japanese.
    s, notes = _strip_jp_parentheticals(english)
    m = re.search(r"\s*:\s*([^A-Za-z]*[\u3040-\u30ff\u3400-\u9fff][^A-Za-z]*)$", s)  # "...counter: ひとつ、ふたつ..."
    if m:
        notes.append(m.group(1).strip())
        s = s[:m.start()]
    if JP_RE.search(s):  # fallback: anything left over, so a spoiler never ships
        notes.append("".join(JP_RE.findall(s)))
        s = JP_RE.sub("", s)
    s = re.sub(r"\s+", " ", s).strip(" +:,;")
    return s, "; ".join(n for n in notes if n)


def sanitize_words(words):
    changed = []
    for w in words:
        clean, note = sanitize_english(w["e"])
        if clean != w["e"]:
            changed.append((w["e"], clean))
            w["e"] = clean
            if note:
                w["n"] = (w["n"] + "; " + note) if w.get("n") else note
    if changed:
        print(f"Sanitised {len(changed)} English meaning(s) containing Japanese:")
        for old, new in changed:
            print(f"  {old!r}  ->  {new!r}")
    assert not any(JP_RE.search(w["e"]) for w in words), "Japanese text survived sanitising"
    return words


# ---------- ids ----------
def assign_ids(words):
    """Stable ids: 'kanji|reading', plus '|english' where that pair collides."""
    counts = Counter((w["k"], w["h"]) for w in words)
    for w in words:
        key = f'{w["k"]}|{w["h"]}'
        if counts[(w["k"], w["h"])] > 1:
            key += f'|{w["e"]}'
        w["id"] = key
    dupes = [k for k, c in Counter(w["id"] for w in words).items() if c > 1]
    if dupes:
        sys.exit(f"Could not make unique word ids; still colliding: {dupes}")
    return words


# ---------- loading ----------
def load_from_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is required to read .xlsx files. Install with:\n"
                  "  pip install openpyxl --break-system-packages")

    wb = openpyxl.load_workbook(path, data_only=True)
    if "N5 Vocabulary" not in wb.sheetnames:
        sys.exit(f"Sheet 'N5 Vocabulary' not found in {path}. Sheets present: {wb.sheetnames}")
    ws = wb["N5 Vocabulary"]

    rows = list(ws.iter_rows(min_row=1, values_only=True))
    header = [str(h).strip() if h else "" for h in rows[0]]
    if header[:len(EXPECTED_HEADER)] != EXPECTED_HEADER:
        print("Warning: header row doesn't match the expected layout.", file=sys.stderr)
        print(f"  Expected: {EXPECTED_HEADER}", file=sys.stderr)
        print(f"  Found:    {header}", file=sys.stderr)

    words, skipped = [], 0
    for r in rows[1:]:
        if not r or not any(r):
            continue
        kanji, hiragana, romaji, english, category, wtype, learned = (list(r) + [None] * 9)[:7]
        if (not hiragana and not kanji) or not english:
            skipped += 1
            continue
        words.append({
            "k": (kanji or "").strip(),
            "h": (hiragana or "").strip(),
            "r": (romaji or "").strip(),
            "e": (english or "").strip(),
            "cat": (category or "General").strip(),
            "type": (wtype or "").strip(),
            "learned": (str(learned).strip().upper() == "YES"),
        })
    if skipped:
        print(f"Note: skipped {skipped} row(s) missing required fields.", file=sys.stderr)
    return assign_ids(sanitize_words(words))


def load_from_json(path):
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    return assign_ids(sanitize_words(data))


# ---------- pulling the live sheet ----------
def _api(url, token, accept):
    req = urllib.request.Request(url, headers={"Authorization": f"token {token}", "Accept": accept,
                                               "User-Agent": "n5-vocab-build"})
    with urllib.request.urlopen(req, timeout=60) as resp:
        return resp.read()


def pull_sheet(repo, sheet_path, token, dest):
    """Download the sheet; return {'date': ISO date, 'sha': short sha} of its last commit."""
    base = f"https://api.github.com/repos/{repo}"
    data = _api(f"{base}/contents/{sheet_path}", token, "application/vnd.github.raw")
    with open(dest, "wb") as f:
        f.write(data)
    info = {"date": None, "sha": None}
    try:
        commits = json.loads(_api(f"{base}/commits?path={sheet_path}&per_page=1", token,
                                  "application/vnd.github.v3+json"))
        if commits:
            info["date"] = commits[0]["commit"]["committer"]["date"][:10]
            info["sha"] = commits[0]["sha"][:7]
    except Exception as e:  # the file itself came through; the label is a nicety
        print(f"Note: couldn't read the sheet's commit date ({e}).", file=sys.stderr)
    return info


# ---------- change report ----------
def report_changes(old, new):
    if not old:
        return
    key = lambda w: f'{w["k"]}|{w["h"]}'   # id-scheme independent, so numeric-id exports compare fine
    om, nm = {key(w): w for w in old}, {key(w): w for w in new}
    added = [k for k in nm if k not in om]
    removed = [k for k in om if k not in nm]
    newly = [k for k in nm if k in om and nm[k]["learned"] and not om[k]["learned"]]
    unlearned = [k for k in nm if k in om and not nm[k]["learned"] and om[k]["learned"]]
    edited = [k for k in nm if k in om and om[k]["e"] != nm[k]["e"]]
    recat = [k for k in nm if k in om and om[k]["cat"] != nm[k]["cat"]]
    show = lambda ks: ", ".join(k.split("|")[0] for k in ks[:8]) + (" ..." if len(ks) > 8 else "")
    print(f"Changes vs previous build: {len(old)} -> {len(new)} words, "
          f"learned {sum(w['learned'] for w in old)} -> {sum(w['learned'] for w in new)}")
    for label, ks in (("added", added), ("removed", removed), ("newly learned", newly),
                      ("no longer learned", unlearned), ("meaning edited", edited), ("category changed", recat)):
        if ks:
            print(f"  {label} ({len(ks)}): {show(ks)}")
    if not any((added, removed, newly, unlearned, edited, recat)):
        print("  (no differences)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", help="Build from this .xlsx (no pull)")
    ap.add_argument("--json", help="Build from an extracted vocab_data.json (no pull)")
    ap.add_argument("--no-pull", action="store_true", help="Use the local .xlsx without pulling the live sheet")
    ap.add_argument("--repo", default=DEFAULT_REPO, help=f"Tutoring repo to pull from (default {DEFAULT_REPO})")
    ap.add_argument("--sheet-path", default=DEFAULT_SHEET_PATH, help="Path of the sheet inside that repo")
    ap.add_argument("--out", default=OUTPUT_PATH, help="Output HTML path (default: index.html)")
    ap.add_argument("--save-json", action="store_true", help="(kept for compatibility; vocab_data.json is always written for sheet builds)")
    args = ap.parse_args()

    sheet_info = {"date": None, "sha": None}
    token = os.environ.get("GITHUB_TOKEN", "").strip()

    if args.json:
        words, src = load_from_json(args.json), args.json
    elif args.xlsx or args.no_pull or not token:
        path = args.xlsx or DEFAULT_XLSX
        if not os.path.exists(path):
            sys.exit(f"No local sheet found ({path}). Set GITHUB_TOKEN to pull the live one, or pass --xlsx / --json.")
        if not args.xlsx and not args.no_pull:
            print("!! GITHUB_TOKEN is not set - building from the LOCAL snapshot, which may be stale. "
                  "Set it to pull the live sheet.", file=sys.stderr)
        words, src = load_from_xlsx(path), path
        mtime = datetime.fromtimestamp(os.path.getmtime(path), timezone.utc).strftime("%Y-%m-%d")
        sheet_info = {"date": None, "sha": None, "local": mtime}
    else:
        try:
            sheet_info = pull_sheet(args.repo, args.sheet_path, token, DEFAULT_XLSX)
        except Exception as e:
            sys.exit(f"Could not pull the live sheet from {args.repo}: {e}\n"
                      "Fix the token/network, or use --no-pull to build from the local snapshot deliberately.")
        words, src = load_from_xlsx(DEFAULT_XLSX), f"{args.repo}:{args.sheet_path}"
        print(f"Pulled live sheet (commit {sheet_info.get('sha')}, {sheet_info.get('date')}).")

    if not words:
        sys.exit(f"No words loaded from {src} - nothing to build.")

    previous = None
    if os.path.exists(DEFAULT_JSON):
        try:
            with open(DEFAULT_JSON, encoding="utf-8") as f:
                previous = json.load(f)
        except Exception:
            previous = None
    report_changes(previous, words)

    if not args.json:
        with open(DEFAULT_JSON, "w", encoding="utf-8") as f:
            json.dump(words, f, ensure_ascii=False, indent=1)
        print(f"Wrote {DEFAULT_JSON} ({len(words)} words)")

    legacy = []
    if os.path.exists(LEGACY_IDS):
        with open(LEGACY_IDS, encoding="utf-8") as f:
            legacy = json.load(f)

    with open(TEMPLATE_PATH, "r", encoding="utf-8") as f:
        template = f.read()

    build_info = json.dumps({
        "builtAt": datetime.now(timezone.utc).strftime("%Y-%m-%d"),
        "count": len(words),
        "sheetDate": sheet_info.get("date") or sheet_info.get("local"),
        "sheetSha": sheet_info.get("sha"),
        "pulled": bool(sheet_info.get("sha")),
    })
    out = template
    for marker, payload in (("VOCAB_JSON", json.dumps(words, ensure_ascii=False)),
                            ("BUILD_INFO", build_info),
                            ("LEGACY_KEYS", json.dumps(legacy, ensure_ascii=False))):
        start, end = f"/*__{marker}__*/", f"/*__END_{marker}__*/"
        i, j = out.find(start), out.find(end)
        if i < 0 or j < 0:
            sys.exit(f"Template placeholder {marker} is missing or was edited.")
        out = out[:i] + start + payload + out[j:]

    with open(args.out, "w", encoding="utf-8") as f:
        f.write(out)
    size_kb = os.path.getsize(args.out) / 1024
    print(f"Built {args.out} - {len(words)} words from {src} ({size_kb:.0f} KB)")


if __name__ == "__main__":
    main()
